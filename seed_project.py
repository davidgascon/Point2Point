#!/usr/bin/env python3
"""
Seed a project into Field Checkout from a P2P workbook.

    pip install openpyxl requests
    python3 seed_project.py \
        --url    https://checkout.yourdomain.com \
        --admin  you@company.com \
        --file   KOKUSAI_P2P_current.xlsm \
        --name   KOKUSAI \
        --job    86259065 \
        --members dylan@co.com jacob@co.com eric@co.com

Reads the Main P2PCO Sheet, creates the project, adds the members, and writes
one point record per row — carrying across every sign-off already on the
sheet, so the app opens showing the real state of the job rather than a blank
list.

Safe to re-run: points are matched on (unit, tag) and updated in place.
"""
from __future__ import annotations

import argparse
import getpass
import re
import sys
import unicodedata

import requests

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl requests")


# --------------------------------------------------------------- sheet read
def norm(v) -> str:
    if v is None:
        return ""
    return " ".join(unicodedata.normalize("NFKD", str(v)).split()).strip().lower()


def cell_state(v):
    """A sign-off cell is initials+date, DNP, N/A, or empty."""
    s = str(v or "").strip()
    if not s:
        return ("", "")
    if re.fullmatch(r"n/?a", s, re.I):
        return ("na", "")
    if s.upper().startswith("DNP"):
        return ("dnp", s)
    return ("done", s)


def read_sheet(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=True)
    ws = wb["Main P2PCO Sheet"] if "Main P2PCO Sheet" in wb.sheetnames else wb.worksheets[0]

    hdr = None
    for r in range(1, min(ws.max_row, 30) + 1):
        row = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "point name" in row and ("unit tag" in row or "type" in row):
            hdr = r
            break
    if not hdr:
        sys.exit("Could not find the header row (Unit Tag / Point Name).")

    headers = {norm(ws.cell(hdr, c).value): c for c in range(1, ws.max_column + 1)}

    def col(*names):
        for n in names:
            if n in headers:
                return headers[n]
            for h, c in headers.items():
                if h.startswith(n):
                    return c
        return None

    C = {
        "unit": col("unit tag"), "tag": col("point name"), "type": col("type"),
        "ctrl": col("controller"), "descr": col("description"),
        "addr": col("address"), "installed": col("installed and labeled"),
        "p2p": col("point to point"), "trends": col("trend/history"),
        "alarm": col("alarm"), "graphic": col("graphic"),
        "startup": col("startup date"), "notes": col("notes"),
    }

    out = []
    for r in range(hdr + 1, ws.max_row + 1):
        tag = ws.cell(r, C["tag"]).value
        if tag is None or not str(tag).strip():
            continue
        if norm(tag) in ("engineering", "electrician", "specialist"):
            continue

        p2p_state, p2p_raw = cell_state(ws.cell(r, C["p2p"]).value if C["p2p"] else None)
        inst_state, inst_raw = cell_state(ws.cell(r, C["installed"]).value if C["installed"] else None)

        checks = {}
        for key, name in (("trends", "Trends"), ("alarm", "Alarm"),
                          ("graphic", "Graphic"), ("startup", "Startup")):
            if not C[key]:
                continue
            st, _ = cell_state(ws.cell(r, C[key]).value)
            if st == "na":
                checks[name] = "na"
            elif st == "done":
                checks[name] = "done"

        rec = {
            "unit": str(ws.cell(r, C["unit"]).value or "").strip() or "Unassigned",
            "tag": str(tag).strip(),
            "type": str(ws.cell(r, C["type"]).value or "-").strip(),
            "descr": str(ws.cell(r, C["descr"]).value or "").strip() if C["descr"] else "",
            "addr": str(ws.cell(r, C["addr"]).value or "").strip() if C["addr"] else "",
            "controller": str(ws.cell(r, C["ctrl"]).value or "").strip() if C["ctrl"] else "",
            "sheet_row": r,
            "checks": checks,
            "notes": str(ws.cell(r, C["notes"]).value or "").strip() if C["notes"] else "",
            "imported_by": "",
            "installed_by": inst_raw if inst_state == "done" else "",
            "p2p_na": p2p_state == "na",
            "status": "",
        }
        if p2p_state == "done":
            rec["status"] = "Pass"
            rec["imported_by"] = p2p_raw
        elif p2p_state == "dnp":
            rec["status"] = "Fail"
            rec["imported_by"] = p2p_raw
        out.append(rec)
    return out


# ------------------------------------------------------------------- client
class PB:
    def __init__(self, url: str):
        self.url = url.rstrip("/")
        self.s = requests.Session()
        self.token = None

    def login_admin(self, email: str, password: str):
        for path in ("/api/collections/_superusers/auth-with-password",
                     "/api/admins/auth-with-password"):
            r = self.s.post(self.url + path,
                            json={"identity": email, "password": password}, timeout=20)
            if r.ok:
                self.token = r.json()["token"]
                self.s.headers["Authorization"] = self.token
                return
        sys.exit(f"Admin login failed: {r.status_code} {r.text[:300]}")

    def find(self, coll: str, filt: str):
        r = self.s.get(f"{self.url}/api/collections/{coll}/records",
                       params={"filter": filt, "perPage": 1}, timeout=20)
        r.raise_for_status()
        items = r.json().get("items", [])
        return items[0] if items else None

    def create(self, coll: str, data: dict):
        r = self.s.post(f"{self.url}/api/collections/{coll}/records", json=data, timeout=30)
        if not r.ok:
            sys.exit(f"Create in {coll} failed: {r.status_code} {r.text[:400]}")
        return r.json()

    def update(self, coll: str, rid: str, data: dict):
        r = self.s.patch(f"{self.url}/api/collections/{coll}/records/{rid}",
                         json=data, timeout=30)
        if not r.ok:
            sys.exit(f"Update in {coll} failed: {r.status_code} {r.text[:400]}")
        return r.json()

    def all(self, coll: str, filt: str = ""):
        out, page = [], 1
        while True:
            r = self.s.get(f"{self.url}/api/collections/{coll}/records",
                           params={"filter": filt, "perPage": 500, "page": page}, timeout=60)
            r.raise_for_status()
            j = r.json()
            out += j["items"]
            if page >= j["totalPages"]:
                return out
            page += 1


# --------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--url", required=True)
    ap.add_argument("--admin", required=True)
    ap.add_argument("--file", required=True)
    ap.add_argument("--name", required=True)
    ap.add_argument("--job", default="")
    ap.add_argument("--location", default="")
    ap.add_argument("--members", nargs="*", default=[])
    a = ap.parse_args()

    pw = getpass.getpass(f"Password for {a.admin}: ")

    rows = read_sheet(a.file)
    units = sorted({r["unit"] for r in rows})
    done = sum(1 for r in rows if r["status"] in ("Pass", "Fail"))
    na = sum(1 for r in rows if r["p2p_na"])
    print(f"Read {len(rows)} points across {len(units)} units "
          f"({done} already signed off, {na} marked N/A)")

    pb = PB(a.url)
    pb.login_admin(a.admin, pw)
    print("Signed in as admin")

    proj = pb.find("projects", f'name="{a.name}"')
    if proj:
        print(f"Project {a.name} exists ({proj['id']}), updating its points")
    else:
        proj = pb.create("projects", {
            "name": a.name, "job": a.job, "location": a.location,
            "status": "active", "settings": {},
        })
        print(f"Created project {proj['id']}")
    pid = proj["id"]

    # ---- members -------------------------------------------------------
    for email in a.members:
        user = pb.find("users", f'email="{email}"')
        if not user:
            print(f"  ! no account for {email} — create it in the dashboard first")
            continue
        if pb.find("project_members", f'project="{pid}" && user="{user["id"]}"'):
            print(f"  = {email} already a member")
            continue
        role = "lead" if user.get("role") in ("lead", "admin") else "tech"
        pb.create("project_members", {"project": pid, "user": user["id"], "role": role})
        print(f"  + {email} added as {role}")

    # ---- points --------------------------------------------------------
    existing = {(p["unit"], p["tag"]): p for p in pb.all("points", f'project="{pid}"')}
    created = updated = 0
    for i, r in enumerate(rows, 1):
        body = dict(r)
        body["project"] = pid
        key = (r["unit"], r["tag"])
        if key in existing:
            pb.update("points", existing[key]["id"], body)
            updated += 1
        else:
            pb.create("points", body)
            created += 1
        if i % 50 == 0:
            print(f"  ... {i}/{len(rows)}")

    print(f"\nDone. {created} points created, {updated} updated.")
    print(f"Project id: {pid}")
    print(f"\nUpload the workbook for exports:")
    print(f"  curl -X POST {a.url}/export/template/{pid} \\")
    print(f'    -H "Authorization: <a lead\'s token>" -F file=@{a.file}')


if __name__ == "__main__":
    main()
